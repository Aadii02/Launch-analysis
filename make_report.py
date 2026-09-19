"""Build a clean Word report from data/launches_clean.csv.

Usage:   uv run python make_report.py
Output:  reports/launch_report.docx   (chart images go to reports/figures/)

Run analysis.ipynb first so data/launches_clean.csv exists.
Re-run this script any time the data changes and the report rebuilds itself.
"""

import textwrap
from datetime import date
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # draw charts to files, no window needed
import matplotlib.pyplot as plt
import pandas as pd
import seaborn as sns
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------- settings you can tweak ----------
DATA = Path("data/launches_clean.csv")
OUT_DIR = Path("reports")
FIG_DIR = OUT_DIR / "figures"
REPORT = OUT_DIR / "launch_report.docx"

TITLE = "Global Launch Analysis"
MIN_LAUNCHES = 20  # rockets with fewer launches are left out of reliability views
TOP_N = 10

NAVY, TEAL, RED, GREY = "#1F3A5F", "#2A9D8F", "#D1495B", "#8D99AE"


# ---------- data ----------
def load():
    df = pd.read_csv(DATA, parse_dates=["net"])
    df["success"] = df["success"].astype(bool)
    return df


def rocket_stats(df):
    g = df.groupby("rocket")["success"].agg(launches="size", successes="sum")
    g["rate"] = g["successes"] / g["launches"] * 100
    return g[g["launches"] >= MIN_LAUNCHES]


def findings(df, stats):
    n = len(df)
    per_year = df["year"].value_counts()
    providers = df["provider"].value_counts()
    rockets = df["rocket"].value_counts()

    out = [
        f"{n:,} completed launches between {int(df['year'].min())} and "
        f"{int(df['year'].max())}, with an overall success rate of "
        f"{df['success'].mean() * 100:.1f}%.",
        f"Busiest year: {int(per_year.idxmax())} with {per_year.max()} launches.",
        f"Most active provider: {providers.index[0]} "
        f"({providers.iloc[0] / n * 100:.0f}% of all launches).",
        f"Most launched rocket: {rockets.index[0]} ({rockets.iloc[0]} launches).",
    ]
    if not stats.empty:
        best, worst = stats["rate"].idxmax(), stats["rate"].idxmin()
        out.append(
            f"Among rockets with at least {MIN_LAUNCHES} launches, {best} has the "
            f"highest success rate ({stats.loc[best, 'rate']:.1f}%) and {worst} "
            f"the lowest ({stats.loc[worst, 'rate']:.1f}%)."
        )
    return out


# ---------- charts ----------
def save(fig, name):
    path = FIG_DIR / name
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)
    return path


def chart_per_year(df):
    counts = df.groupby(["year", "success"]).size().unstack(fill_value=0)
    counts = counts.reindex(columns=[True, False], fill_value=0)
    fig, ax = plt.subplots(figsize=(8, 3.8))
    ax.bar(counts.index, counts[True], color=TEAL, label="Success")
    ax.bar(counts.index, counts[False], bottom=counts[True], color=RED,
           label="Failure / partial failure")
    ax.set_xlabel("Year")
    ax.set_ylabel("Launches")
    ax.legend(frameon=False)
    return save(fig, "launches_per_year.png")


def chart_by_decade(df):
    d = df.assign(decade=(df["year"] // 10) * 10).groupby("decade")["success"].agg(
        ["mean", "size"]
    )
    fig, ax = plt.subplots(figsize=(8, 3.8))
    ax.bar([f"{int(x)}s" for x in d.index], d["mean"] * 100, color=NAVY)
    for i, (rate, n) in enumerate(zip(d["mean"], d["size"])):
        ax.text(i, rate * 100 + 1.5, f"n={n}", ha="center", fontsize=8, color=GREY)
    ax.set_ylim(0, 108)
    ax.set_ylabel("Success rate (%)")
    return save(fig, "success_by_decade.png")


def chart_rocket_reliability(top):
    top = top.sort_values("rate")
    labels = [f"{name}  (n={int(n)})" for name, n in zip(top.index, top["launches"])]
    fig, ax = plt.subplots(figsize=(8, 4.6))
    ax.scatter(top["rate"], labels, color=NAVY, s=45, zorder=3)
    ax.set_xlim(max(0, top["rate"].min() - 5), 100.5)
    ax.set_xlabel("Success rate (%)")
    return save(fig, "rocket_reliability.png")


def chart_top_counts(df, col, fname, exclude=("Unknown",)):
    counts = df[~df[col].isin(exclude)][col].value_counts().head(TOP_N).sort_values()
    fig, ax = plt.subplots(figsize=(8, 4))
    labels = [textwrap.shorten(str(x), width=30, placeholder="...") for x in counts.index]
    ax.barh(labels, counts.values, color=NAVY)
    ax.set_xlabel("Launches")
    return save(fig, fname)


# ---------- docx helpers ----------
def set_font(style, name, size=None, color=None, bold=None):
    style.font.name = name
    rfonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
    if rfonts is not None:  # drop theme fonts so our font actually applies
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
    if size:
        style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    if bold is not None:
        style.font.bold = bold


def shade(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shd)


def light_borders(table, color="D9DEE5"):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GREY.lstrip("#"))


def add_figure(doc, path, caption, number):
    doc.add_picture(str(path), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Figure {number}. {caption}")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GREY.lstrip("#"))


def add_rocket_table(doc, top):
    rows = top.sort_values("launches", ascending=False)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    light_borders(table)
    widths = [Cm(7), Cm(3), Cm(3), Cm(3)]

    for i, text in enumerate(["Rocket", "Launches", "Successes", "Success rate"]):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].alignment = (
            WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        )
        shade(cell, NAVY)

    for r, (name, row) in enumerate(rows.iterrows()):
        cells = table.add_row().cells
        values = [str(name), f"{int(row['launches'])}", f"{int(row['successes'])}",
                  f"{row['rate']:.1f}%"]
        for i, text in enumerate(values):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(text)
            run.font.size = Pt(10)
            cells[i].paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            )
            if r % 2 == 1:
                shade(cells[i], "F3F5F8")

    for i, w in enumerate(widths):
        table.columns[i].width = w
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w
            para = row.cells[i].paragraphs[0]
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(3)


# ---------- build the report ----------
def build_doc(df, stats, figs):
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)

    set_font(doc.styles["Normal"], "Calibri", size=11)
    set_font(doc.styles["Heading 1"], "Calibri", size=16, color=NAVY, bold=True)
    set_font(doc.styles["Heading 2"], "Calibri", size=13, color=TEAL, bold=True)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)

    # title block
    p = doc.add_paragraph()
    run = p.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string(NAVY.lstrip("#"))
    p = doc.add_paragraph()
    run = p.add_run(
        f"Launches {int(df['year'].min())} to {int(df['year'].max())}  |  "
        f"Generated {date.today():%d %B %Y}"
    )
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(GREY.lstrip("#"))

    doc.add_heading("Key findings", level=1)
    for line in findings(df, stats):
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Launches over time", level=1)
    add_figure(doc, figs["per_year"], "Launches per year, split by outcome.", 1)

    doc.add_heading("Reliability", level=1)
    add_figure(doc, figs["decade"],
               "Success rate by decade (n = launches in that decade).", 2)

    if not stats.empty:
        top = stats.sort_values("launches", ascending=False).head(15)
        doc.add_heading("Reliability by rocket", level=2)
        add_figure(doc, figs["rocket"],
                   f"Success rate of the most-launched rockets (at least {MIN_LAUNCHES} launches).",
                   3)
        add_rocket_table(doc, top)

    doc.add_heading("Providers and orbits", level=1)
    add_figure(doc, figs["providers"], f"Top {TOP_N} providers by launch count.", 4)
    add_figure(doc, figs["orbits"],
               f"Top {TOP_N} target orbits (launches with unknown orbit excluded).", 5)

    doc.add_heading("About the data", level=1)
    notes = [
        "Source: Launch Library 2 by The Space Devs, pulled with fetch_data.py.",
        "Only launches with a recorded outcome are included; scheduled or unknown launches are dropped.",
        "A launch counts as a success only if its status is Launch Successful. Partial failures count as failures.",
        "Rocket, provider, and orbit names are shown as they appear in the source data.",
    ]
    for line in notes:
        doc.add_paragraph(line, style="List Bullet")

    doc.save(REPORT)


def main():
    if not DATA.exists():
        raise SystemExit(f"{DATA} not found. Run analysis.ipynb first to create it.")

    FIG_DIR.mkdir(parents=True, exist_ok=True)
    sns.set_theme(style="whitegrid",
                  rc={"axes.spines.top": False, "axes.spines.right": False})

    df = load()
    stats = rocket_stats(df)

    figs = {
        "per_year": chart_per_year(df),
        "decade": chart_by_decade(df),
        "providers": chart_top_counts(df, "provider", "top_providers.png"),
        "orbits": chart_top_counts(df, "orbit_abbrev", "top_orbits.png"),
    }
    if not stats.empty:
        top = stats.sort_values("launches", ascending=False).head(15)
        figs["rocket"] = chart_rocket_reliability(top)

    build_doc(df, stats, figs)
    print(f"Report written to {REPORT}")


if __name__ == "__main__":
    main()

